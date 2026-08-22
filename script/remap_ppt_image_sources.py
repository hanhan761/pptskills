#!/usr/bin/env python3
"""Remap image-source records to a revised deck by embedded byte identity."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from template_ppt import find_repo_root, read_json, resolve_repo_path


def digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def dhash(path: Path, size: int = 16) -> str:
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("L").resize(
            (size + 1, size), Image.Resampling.LANCZOS
        )
        pixels = list(image.getdata())
    value = 0
    for row in range(size):
        offset = row * (size + 1)
        for column in range(size):
            value = (value << 1) | int(
                pixels[offset + column] > pixels[offset + column + 1]
            )
    return f"{value:0{size * size // 4}x}"


def picture_inventory(deck_path: Path) -> list[dict[str, Any]]:
    deck = Presentation(deck_path)
    records: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(deck.slides, 1):
        occurrences: dict[str, int] = {}
        for shape in slide.shapes:
            name = str(shape.name)
            occurrences[name] = occurrences.get(name, 0) + 1
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            records.append(
                {
                    "slide": slide_number,
                    "shape": name,
                    "occurrence": occurrences[name],
                    "sha256": digest(shape.image.blob),
                }
            )
    return records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--existing-sources", required=True)
    parser.add_argument("--additional-map", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    repo = find_repo_root()
    deck_path = resolve_repo_path(args.input, repo)
    existing_path = resolve_repo_path(args.existing_sources, repo)
    additional_path = resolve_repo_path(args.additional_map, repo)
    output_path = resolve_repo_path(args.output, repo)
    existing = read_json(existing_path)
    additional = read_json(additional_path).get("replacements", [])
    pictures = picture_inventory(deck_path)
    claimed: set[tuple[int, str, int]] = set()
    errors: list[str] = []
    remapped: list[dict[str, Any]] = []

    def match(file_value: str, shape_name: str, occurrence: int) -> dict[str, Any] | None:
        source_file = resolve_repo_path(file_value, repo)
        if not source_file.is_file():
            errors.append(f"source file not found: {source_file}")
            return None
        file_sha = digest(source_file.read_bytes())
        candidates = [
            item
            for item in pictures
            if item["shape"] == shape_name
            and int(item["occurrence"]) == occurrence
            and item["sha256"] == file_sha
            and (item["slide"], item["shape"], item["occurrence"]) not in claimed
        ]
        if len(candidates) != 1:
            errors.append(
                f"expected one current picture for {shape_name!r} occurrence {occurrence} "
                f"and {source_file}; found {len(candidates)}"
            )
            return None
        return candidates[0]

    for record in existing.get("images", []):
        occurrence = int(record.get("occurrence", 1))
        item = match(str(record["file"]), str(record["shape"]), occurrence)
        if item is None:
            continue
        updated = dict(record)
        updated["original_slide"] = int(record["slide"])
        updated["slide"] = int(item["slide"])
        updated["occurrence"] = occurrence
        remapped.append(updated)
        claimed.add((item["slide"], item["shape"], item["occurrence"]))

    for record in additional:
        occurrence = int(record.get("occurrence", 1))
        file_value = str(record["image"])
        item = match(file_value, str(record["shape"]), occurrence)
        if item is None:
            continue
        source_file = resolve_repo_path(file_value, repo)
        remapped.append(
            {
                "slide": int(item["slide"]),
                "shape": str(record["shape"]),
                "occurrence": occurrence,
                "file": file_value,
                "caption": str(record["caption"]),
                "credit": str(record["credit"]),
                "source_page": str(record["source_page"]),
                "direct_url": str(record.get("direct_url", record["source_page"])),
                "sha256": digest(source_file.read_bytes()),
                "dhash16": dhash(source_file),
                "verified": True,
            }
        )
        claimed.add((item["slide"], item["shape"], item["occurrence"]))

    unclaimed = [
        item
        for item in pictures
        if (item["slide"], item["shape"], item["occurrence"]) not in claimed
    ]
    allowed_fixed = [
        item
        for item in unclaimed
        if int(item["slide"]) in {1, len(Presentation(deck_path).slides)}
        and str(item["shape"]) in {"图片 1", "Picture 8"}
    ]
    unexpected = [item for item in unclaimed if item not in allowed_fixed]
    if unexpected:
        errors.append(f"unclaimed non-family picture shapes remain: {unexpected}")
    if len(allowed_fixed) != 4:
        errors.append(
            f"expected four cover/closing family picture assets; found {len(allowed_fixed)}"
        )

    output = {
        "policy_sources": existing.get("policy_sources", []),
        "images": sorted(
            remapped,
            key=lambda item: (
                int(item["slide"]),
                str(item["shape"]),
                int(item.get("occurrence", 1)),
            ),
        ),
        "excluded_fixed_family_assets": allowed_fixed,
        "remap": {
            "deck": str(deck_path),
            "picture_shapes": len(pictures),
            "sourced_picture_shapes": len(remapped),
            "fixed_family_picture_shapes": len(allowed_fixed),
        },
        "errors": errors,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(output["remap"] | {"errors": errors}, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
