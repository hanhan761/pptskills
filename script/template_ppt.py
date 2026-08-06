#!/usr/bin/env python3
"""Mechanical tooling for a template-first PowerPoint workflow.

This script copies and inspects existing slides. It never creates slide layouts.
"""

from __future__ import annotations

import argparse
import copy
import gc
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import time
import zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Iterator
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


CATEGORIES = ("封面", "目录", "内容")
PPTX_FORMAT = 24
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
IMAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)
CONTENT_TYPES = {
    ".bmp": "image/bmp",
    ".gif": "image/gif",
    ".jpeg": "image/jpeg",
    ".jpg": "image/jpeg",
    ".png": "image/png",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}

for prefix, uri in NS.items():
    if prefix not in {"rel", "ct"}:
        ET.register_namespace(prefix, uri)


def fail(message: str) -> None:
    raise SystemExit(f"ERROR: {message}")


def read_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(f"JSON file not found: {path}")
    except json.JSONDecodeError as exc:
        fail(f"Invalid JSON in {path}: {exc}")


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_repo_root(start: Path | None = None) -> Path:
    cursor = (start or Path.cwd()).resolve()
    if cursor.is_file():
        cursor = cursor.parent
    for candidate in (cursor, *cursor.parents):
        if (candidate / "AGENTS.md").is_file() and (candidate / "模板").is_dir():
            return candidate
    fail("Could not find an ancestor containing AGENTS.md and 模板/")


def resolve_repo_path(value: str | Path, repo: Path) -> Path:
    path = Path(value)
    return (repo / path).resolve() if not path.is_absolute() else path.resolve()


def relative_or_absolute(path: Path, repo: Path) -> str:
    try:
        return path.resolve().relative_to(repo.resolve()).as_posix()
    except ValueError:
        return str(path.resolve())


def safe_name(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" ._")
    return cleaned or "template"


def shape_record(shape: Any) -> dict[str, Any]:
    record: dict[str, Any] = {
        "name": shape.name,
        "type": str(shape.shape_type),
        "left": int(shape.left),
        "top": int(shape.top),
        "width": int(shape.width),
        "height": int(shape.height),
        "rotation": float(shape.rotation),
    }
    if getattr(shape, "has_text_frame", False):
        record["paragraphs"] = [paragraph.text for paragraph in shape.text_frame.paragraphs]
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        record["children"] = [shape_record(child) for child in shape.shapes]
    return record


def constraints_path(template_path: Path) -> Path:
    return template_path.with_suffix(template_path.suffix + ".constraints.json")


def validate_constraints_shape(data: dict[str, Any], path: Path) -> list[str]:
    errors: list[str] = []
    slots = data.get("text_slots")
    if data.get("schema_version") != 1:
        errors.append(f"{path}: schema_version must be 1")
    if not isinstance(slots, list):
        errors.append(f"{path}: text_slots must be a list")
        return errors
    if not slots:
        if data.get("image_only") is not True:
            errors.append(
                f"{path}: an empty text_slots list requires image_only: true"
            )
        return errors
    seen: set[tuple[str, int | None]] = set()
    for number, slot in enumerate(slots, 1):
        if not isinstance(slot, dict) or not slot.get("shape"):
            errors.append(f"{path}: text slot {number} is missing shape")
            continue
        occurrence = int(slot["occurrence"]) if "occurrence" in slot else None
        key = (str(slot["shape"]), occurrence)
        if key in seen:
            errors.append(f"{path}: duplicate text slot selector {key}")
        seen.add(key)
        minimum = int(slot.get("min_chars", 0))
        maximum = int(slot.get("max_chars", 0))
        if minimum < 0 or maximum < minimum or maximum == 0:
            errors.append(f"{path}: invalid character range for {key}")
        if "line_count" in slot:
            line_count = int(slot["line_count"])
            if line_count < 1:
                errors.append(f"{path}: line_count must be positive for {key}")
            for field in ("line_min_chars", "line_max_chars"):
                values = slot.get(field)
                if values is not None and (
                    not isinstance(values, list) or len(values) != line_count
                ):
                    errors.append(
                        f"{path}: {field} must contain {line_count} values for {key}"
                    )
        style_roles = slot.get("style_roles")
        if style_roles is not None:
            if not isinstance(style_roles, dict) or not style_roles:
                errors.append(f"{path}: style_roles must be a non-empty object for {key}")
            else:
                for role, source_run in style_roles.items():
                    if not isinstance(role, str) or not role.strip():
                        errors.append(f"{path}: style role names must be non-empty for {key}")
                    if not isinstance(source_run, int) or isinstance(source_run, bool) or source_run < 1:
                        errors.append(
                            f"{path}: style role '{role}' must reference a positive source run for {key}"
                        )
            required_roles = slot.get("required_style_roles", [])
            if not isinstance(required_roles, list) or not all(
                isinstance(role, str) and role for role in required_roles
            ):
                errors.append(
                    f"{path}: required_style_roles must be a list of non-empty strings for {key}"
                )
            elif isinstance(style_roles, dict):
                unknown = [role for role in required_roles if role not in style_roles]
                if unknown:
                    errors.append(
                        f"{path}: required_style_roles contains unknown roles {unknown} for {key}"
                    )
    return errors


def slide_record(slide: Any, number: int) -> dict[str, Any]:
    shapes = [shape_record(shape) for shape in slide.shapes]
    return {
        "slide": number,
        "shape_count": len(shapes),
        "picture_count": sum(
            1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ),
        "shapes": shapes,
    }


def open_presentation(path: Path) -> Presentation:
    if not path.is_file():
        fail(f"PPTX not found: {path}")
    try:
        return Presentation(str(path))
    except Exception as exc:
        fail(f"Cannot read PPTX {path}: {exc}")


def cmd_index(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    library = repo / "模板"
    entries: list[dict[str, Any]] = []
    errors: list[str] = []
    ledger_path = library / "学习记录.json"
    if ledger_path.is_file():
        ledger = read_json(ledger_path)
        curation_status = ledger.get("validation", {}).get("curation_status")
        if curation_status == "pending":
            errors.append(
                "模板学习仍处于待策展状态；必须先逐页完成 keep/merge/discard "
                "并运行 script/curate_template_library.py --apply"
            )
    for category in CATEGORIES:
        folder = library / category
        for path in sorted(folder.rglob("*.pptx")):
            prs = open_presentation(path)
            count = len(prs.slides)
            if count != 1:
                errors.append(f"{relative_or_absolute(path, repo)} has {count} slides; expected 1")
            contract_path = constraints_path(path)
            contract_data: dict[str, Any] | None = None
            if not contract_path.is_file():
                errors.append(
                    f"{relative_or_absolute(path, repo)} is missing its capacity contract"
                )
            else:
                try:
                    contract_data = json.loads(contract_path.read_text(encoding="utf-8"))
                    errors.extend(validate_constraints_shape(contract_data, contract_path))
                except (OSError, json.JSONDecodeError) as exc:
                    errors.append(f"{contract_path}: cannot read capacity contract: {exc}")
            entries.append(
                {
                    "path": relative_or_absolute(path, repo),
                    "category": category,
                    "slide_count": count,
                    "slide_width": int(prs.slide_width),
                    "slide_height": int(prs.slide_height),
                    "sha256": sha256(path),
                    "constraints": relative_or_absolute(contract_path, repo)
                    if contract_data is not None
                    else None,
                    "slides": [
                        slide_record(slide, number)
                        for number, slide in enumerate(prs.slides, 1)
                    ],
                }
            )
    result = {"schema_version": 1, "templates": entries, "errors": errors}
    if args.write:
        output = library / "索引.json"
        write_json(output, result)
        print(f"Wrote {output} with {len(entries)} template(s)")
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(2)


def cmd_inspect(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    path = resolve_repo_path(args.input, repo)
    prs = open_presentation(path)
    result = {
        "path": relative_or_absolute(path, repo),
        "slide_width": int(prs.slide_width),
        "slide_height": int(prs.slide_height),
        "slides": [
            slide_record(slide, number) for number, slide in enumerate(prs.slides, 1)
        ],
    }
    if args.output:
        write_json(resolve_repo_path(args.output, repo), result)
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))


@contextmanager
def powerpoint_app() -> Iterator[Any]:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        fail("PowerPoint automation requires pywin32 on Windows")
    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        app.DisplayAlerts = 0
        yield app
    except Exception as exc:
        fail(f"PowerPoint automation failed: {exc}")
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        app = None
        pythoncom.CoUninitialize()
        gc.collect()
        # Office may return from Quit before the out-of-process COM server has
        # completely released RPC registration.  A short settle interval keeps
        # batched long-deck composition from racing the next DispatchEx call.
        time.sleep(0.5)


def copy_with_retry(source: Path, destination: Path) -> None:
    last_error: PermissionError | None = None
    for attempt in range(12):
        try:
            shutil.copy2(source, destination)
            return
        except PermissionError as exc:
            last_error = exc
            time.sleep(0.25 * (attempt + 1))
    if last_error is not None:
        raise last_error


def extract_one(app: Any, source: Path, slide_number: int, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ppt-extract-") as temporary:
        temporary_source = Path(temporary) / "source.pptx"
        temporary_output = Path(temporary) / "slide.pptx"
        shutil.copy2(source, temporary_source)
        presentation = app.Presentations.Open(
            str(temporary_source), False, False, False
        )
        try:
            for index in range(int(presentation.Slides.Count), 0, -1):
                if index != slide_number:
                    presentation.Slides(index).Delete()
            if presentation.Slides.Count != 1:
                fail(f"Extraction produced {presentation.Slides.Count} slides for {output}")
            presentation.SaveAs(str(temporary_output), PPTX_FORMAT)
        finally:
            try:
                presentation.Close()
            except Exception:
                if not temporary_output.is_file():
                    raise
        copy_with_retry(temporary_output, output)


def cmd_extract(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    source = resolve_repo_path(args.source, repo)
    if not source.is_file():
        fail(f"Source PPTX not found: {source}")
    if args.output_dir:
        output_dir = resolve_repo_path(args.output_dir, repo)
    else:
        if args.category not in CATEGORIES:
            fail(f"--category must be one of: {', '.join(CATEGORIES)}")
        output_dir = repo / "模板" / args.category
    names = args.names or []
    if names and len(names) != len(args.slides):
        fail("--names must contain exactly one name for every requested slide")
    outputs: list[Path] = []
    count = len(open_presentation(source).slides)
    with powerpoint_app() as app:
        for position, slide_number in enumerate(args.slides):
            if slide_number < 1 or slide_number > count:
                fail(f"Slide {slide_number} is outside 1..{count}")
            stem = safe_name(names[position]) if names else f"{safe_name(source.stem)}__s{slide_number:03d}"
            output = output_dir / f"{stem}.pptx"
            if output.exists() and not args.overwrite:
                fail(f"Output already exists; pass --overwrite to replace it: {output}")
            extract_one(app, source, slide_number, output)
            outputs.append(output)
    for output in outputs:
        print(relative_or_absolute(output, repo))


def iter_pptx_shapes(shapes: Any) -> Iterator[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_pptx_shapes(shape.shapes)
        else:
            yield shape


def pptx_shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return str(shape.text_frame.text or "")
    if getattr(shape, "has_table", False):
        return "\n".join(
            str(cell.text or "")
            for row in shape.table.rows
            for cell in row.cells
        )
    return ""


def pptx_shape_font_size(shape: Any) -> float:
    sizes: list[float] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(float(run.font.size.pt))
    return max(sizes, default=0.0)


def compact_text(value: str) -> str:
    return re.sub(r"\s+", "", value.replace("\x0b", " ").replace("\r", " "))


def display_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\x0b", " ").replace("\r", " ")).strip()


def inferred_topic(slide: Any, slide_width: int, slide_height: int) -> str:
    candidates: list[tuple[float, str]] = []
    for position, shape in enumerate(iter_pptx_shapes(slide.shapes)):
        text = display_text(pptx_shape_text(shape))
        visible = compact_text(text)
        if len(visible) < 2 or re.fullmatch(r"[\d./—–-]+", visible):
            continue
        font_size = pptx_shape_font_size(shape)
        top_ratio = max(0.0, min(1.0, float(shape.top) / max(1, slide_height)))
        score = font_size * 2.0 + (1.0 - top_ratio) * 55.0 - position * 0.05
        if 4 <= len(visible) <= 36:
            score += 25.0
        elif len(visible) > 70:
            score -= 45.0
        if "标题" in str(shape.name):
            score += 30.0
        if any(token in visible for token in ("示例组织", "汇报人：", "2026年")):
            score -= 80.0
        candidates.append((score, text))
    if not candidates:
        return "无标题页面"
    chosen = max(candidates, key=lambda item: item[0])[1]
    chosen = re.sub(r"^[0-9一二三四五六七八九十.、\-]+", "", chosen).strip()
    return compact_text(chosen)[:24] or "无标题页面"


def source_family_and_variant(source: Path) -> tuple[str, str]:
    stem = source.stem
    if "轻量化构件" in stem:
        return "通用蓝白_示例来源", "v4"
    if "v4-ws2" in stem.lower():
        return "通用蓝白_示例来源", "ws2"
    if "v3" in stem.lower():
        return "通用蓝白_示例来源", "v3"
    return "通用蓝白_示例来源", "答辩版"


def infer_category(slide_number: int, texts: list[str]) -> str:
    combined = "".join(compact_text(text) for text in texts)
    if slide_number == 1 or any(
        marker in combined for marker in ("谢谢大家", "谢谢敬请批评指正", "恳请各位专家批评指正")
    ):
        return "封面"
    if any(marker in combined for marker in ("汇报提纲", "答辩提纲", "主要内容")):
        return "目录"
    return "内容"


def infer_structure(category: str, topic: str, slide: Any) -> str:
    shapes = list(iter_pptx_shapes(slide.shapes))
    pictures = sum(1 for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    tables = sum(1 for shape in shapes if getattr(shape, "has_table", False))
    text_shapes = sum(1 for shape in shapes if compact_text(pptx_shape_text(shape)))
    if category == "封面":
        if any(marker in topic for marker in ("谢谢", "批评指正")):
            return "致谢结束页"
        return "项目答辩封面"
    if category == "目录":
        return "章节导航"
    if any(marker in topic for marker in ("经费", "预算", "资金配置")):
        return "预算表"
    if any(marker in topic for marker in ("进度", "计划")) and tables:
        return "甘特时间轴"
    if any(marker in topic for marker in ("团队", "工作基础")) and pictures >= 3:
        return "团队与基础"
    if any(marker in topic for marker in ("指标", "目标")):
        return "指标卡片"
    if any(marker in topic for marker in ("体系", "总体方案", "闭环", "路径")):
        return "体系流程图"
    if any(marker in topic for marker in ("平台", "原型", "软件", "系统")) and pictures:
        return "系统界面"
    if tables:
        return "数据表格"
    if pictures >= 6:
        return "多图矩阵"
    if pictures >= 3:
        return "多图卡片"
    if pictures:
        return "图文组合"
    if text_shapes >= 10:
        return "多模块图解"
    return "文字图解"


def com_rendered_lines(text_range: Any) -> list[str]:
    lines: list[str] = []
    seen: set[tuple[int, int]] = set()
    maximum = min(max(int(text_range.Length) + 2, 4), 250)
    for number in range(1, maximum):
        try:
            line = text_range.Lines(number, 1)
            key = (int(line.Start), int(line.Length))
            if key in seen or key[1] <= 0:
                break
            seen.add(key)
            lines.append(str(line.Text or ""))
        except Exception:
            break
    return lines


def com_font_size(text_range: Any) -> float:
    try:
        value = float(text_range.Font.Size)
        if value > 0:
            return value
    except Exception:
        pass
    length = max(1, int(getattr(text_range, "Length", 1)))
    positions = sorted({1, max(1, length // 2), length})
    values: list[float] = []
    for position in positions:
        try:
            value = float(text_range.Characters(position, 1).Font.Size)
            if value > 0:
                values.append(value)
        except Exception:
            pass
    return max(values, default=0.0)


def com_text_record(shape: Any, occurrence: int) -> dict[str, Any] | None:
    try:
        if bool(shape.HasTable):
            texts: list[str] = []
            sizes: list[float] = []
            for row in range(1, int(shape.Table.Rows.Count) + 1):
                for column in range(1, int(shape.Table.Columns.Count) + 1):
                    text_range = shape.Table.Cell(row, column).Shape.TextFrame.TextRange
                    texts.append(str(text_range.Text or ""))
                    size = com_font_size(text_range)
                    if size:
                        sizes.append(size)
            text = "\n".join(texts)
            if not compact_text(text):
                return None
            return {
                "shape": str(shape.Name),
                "occurrence": occurrence,
                "text": text,
                "lines": [],
                "font_size": max(sizes, default=0.0),
                "top": float(shape.Top),
                "height": float(shape.Height),
                "table": True,
            }
    except Exception:
        pass
    try:
        if not bool(shape.HasTextFrame) or not bool(shape.TextFrame.HasText):
            return None
        text_range = shape.TextFrame.TextRange
        text = str(text_range.Text or "")
        if not compact_text(text):
            return None
        return {
            "shape": str(shape.Name),
            "occurrence": occurrence,
            "text": text,
            "lines": com_rendered_lines(text_range),
            "font_size": com_font_size(text_range),
            "top": float(shape.Top),
            "height": float(shape.Height),
            "table": False,
        }
    except Exception:
        return None


def collect_com_text_records(slide: Any) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    seen: dict[str, int] = {}

    def visit(collection: Any) -> None:
        for index in range(1, int(collection.Count) + 1):
            shape = collection.Item(index)
            name = str(shape.Name)
            seen[name] = seen.get(name, 0) + 1
            try:
                is_group = int(shape.Type) == int(MSO_SHAPE_TYPE.GROUP)
            except Exception:
                is_group = False
            if is_group:
                visit(shape.GroupItems)
                continue
            record = com_text_record(shape, seen[name])
            if record is not None:
                records.append(record)

    visit(slide.Shapes)
    return records


def run_style_signature(run: ET.Element) -> tuple[str, ...]:
    properties = run.find("./a:rPr", NS)
    if properties is None:
        return ("default",)
    fill = properties.find("./a:solidFill", NS)
    latin = properties.find("./a:latin", NS)
    east_asian = properties.find("./a:ea", NS)
    return (
        str(properties.get("b", "")),
        str(properties.get("i", "")),
        str(properties.get("sz", "")),
        ET.tostring(fill, encoding="unicode") if fill is not None else "",
        str(latin.get("typeface", "")) if latin is not None else "",
        str(east_asian.get("typeface", "")) if east_asian is not None else "",
    )


def xml_container_style_roles(container: ET.Element) -> tuple[int, dict[str, int], list[str]]:
    paragraphs = container.findall(".//a:p", NS)
    if len(paragraphs) != 1:
        return len(paragraphs), {}, []
    run_tag = f"{{{NS['a']}}}r"
    field_tag = f"{{{NS['a']}}}fld"
    runs = [
        child
        for child in list(paragraphs[0])
        if child.tag in {run_tag, field_tag}
        and child.find(".//a:t", NS) is not None
        and compact_text(child.find(".//a:t", NS).text or "")
    ]
    signatures: list[tuple[str, ...]] = []
    first_indexes: list[int] = []
    for index, run in enumerate(runs, 1):
        signature = run_style_signature(run)
        if signature not in signatures:
            signatures.append(signature)
            first_indexes.append(index)
    if len(signatures) < 2 or len(signatures) > 4:
        return 1, {}, []
    roles: dict[str, int] = {"base": first_indexes[0]}
    for position, source_run in enumerate(first_indexes[1:], 1):
        role = "accent" if position == 1 else f"accent{position}"
        roles[role] = source_run
    return 1, roles, list(roles)


def capacity_range(count: int) -> tuple[int, int]:
    if count <= 4:
        return max(1, count - 1), count + 2
    if count <= 12:
        return max(1, math.floor(count * 0.65)), max(count, math.ceil(count * 1.30))
    if count <= 40:
        return max(1, math.floor(count * 0.70)), max(count, math.ceil(count * 1.20))
    if count <= 120:
        return max(1, math.floor(count * 0.72)), max(count, math.ceil(count * 1.16))
    return max(1, math.floor(count * 0.70)), max(count, math.ceil(count * 1.12))


def should_fix_rendered_lines(
    record: dict[str, Any], slide_height: float, paragraph_count: int
) -> bool:
    count = visible_char_count(str(record["text"]))
    lines = record["lines"]
    if record["table"] or paragraph_count != 1 or not 1 <= len(lines) <= 3 or count > 80:
        return False
    top_ratio = float(record["top"]) / max(1.0, slide_height)
    height_ratio = float(record["height"]) / max(1.0, slide_height)
    return (
        float(record["font_size"]) >= 18.0
        or top_ratio <= 0.22
        or (len(lines) == 1 and height_ratio <= 0.08 and count <= 60)
    )


def capacity_contract_from_slide(
    slide: Any,
    source: Path,
    slide_number: int,
    repo: Path,
    slide_height: float,
) -> dict[str, Any]:
    with zipfile.ZipFile(source, "r") as archive:
        root = ET.fromstring(archive.read(slide_xml_name(slide_number)))
    xml_name_counts: dict[str, int] = {}
    for node in root.findall(".//p:cNvPr", NS):
        name = str(node.get("name", ""))
        xml_name_counts[name] = xml_name_counts.get(name, 0) + 1
    slots: list[dict[str, Any]] = []
    for record in collect_com_text_records(slide):
        name = str(record["shape"])
        if "灯片编号" in name:
            continue
        occurrence = int(record["occurrence"]) if xml_name_counts.get(name, 0) > 1 else None
        try:
            container = named_shape_container(root, name, occurrence)
        except SystemExit:
            continue
        paragraphs = container.findall(".//a:p", NS)
        paragraph_count = len(paragraphs)
        count = visible_char_count(str(record["text"]))
        minimum, maximum = capacity_range(count)
        slot: dict[str, Any] = {
            "shape": name,
            "required": True,
            "min_chars": minimum,
            "max_chars": maximum,
        }
        if occurrence is not None:
            slot["occurrence"] = occurrence
        if should_fix_rendered_lines(record, slide_height, paragraph_count):
            line_counts = [
                visible_char_count(line) for line in record["lines"] if compact_text(line)
            ]
            if line_counts:
                line_mins = [max(1, math.floor(value * 0.65)) for value in line_counts]
                line_maxes = [max(value, math.ceil(value * 1.20)) for value in line_counts]
                slot["line_count"] = len(line_counts)
                slot["line_min_chars"] = line_mins
                slot["line_max_chars"] = line_maxes
                slot["min_chars"] = sum(line_mins)
                slot["max_chars"] = sum(line_maxes)
                _, roles, required_roles = xml_container_style_roles(container)
                if roles:
                    slot["style_roles"] = roles
                    slot["required_style_roles"] = required_roles
        slots.append(slot)
    contract: dict[str, Any] = {
        "schema_version": 1,
        "source": {
            "deck": relative_or_absolute(source, repo),
            "slide": slide_number,
            "sha256": sha256(source),
        },
        "text_slots": slots,
    }
    if not slots:
        contract["image_only"] = True
    return contract


def cmd_learn(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    if args.overwrite and args.resume:
        fail("--overwrite and --resume are mutually exclusive")
    source_dir = resolve_repo_path(args.source_dir, repo)
    sources = sorted(source_dir.glob("*.pptx"))
    if not sources:
        fail(f"No PPTX files found in {source_dir}")
    ledger: dict[str, Any] = {"schema_version": 1, "sources": []}
    for source in sources:
        prs = open_presentation(source)
        family, variant = source_family_and_variant(source)
        source_entry: dict[str, Any] = {
            "source": relative_or_absolute(source, repo),
            "sha256": sha256(source),
            "slide_count": len(prs.slides),
            "slides": [],
        }
        with powerpoint_app() as inspection_app:
            source_com = inspection_app.Presentations.Open(
                str(source), True, False, False
            )
            try:
                slide_height = float(source_com.PageSetup.SlideHeight)
                contracts = [
                    capacity_contract_from_slide(
                        source_com.Slides.Item(slide_number),
                        source,
                        slide_number,
                        repo,
                        slide_height,
                    )
                    for slide_number in range(1, len(prs.slides) + 1)
                ]
            finally:
                source_com.Close()
        for slide_number, pptx_slide in enumerate(prs.slides, 1):
            texts = [
                pptx_shape_text(shape)
                for shape in iter_pptx_shapes(pptx_slide.shapes)
                if compact_text(pptx_shape_text(shape))
            ]
            category = infer_category(slide_number, texts)
            topic = inferred_topic(
                pptx_slide, int(prs.slide_width), int(prs.slide_height)
            )
            structure = infer_structure(category, topic, pptx_slide)
            stem = safe_name(
                f"{family}_{structure}_{topic}_{variant}_P{slide_number:02d}"
            )
            output = repo / "模板" / category / f"{stem}.pptx"
            contract_path = constraints_path(output)
            if output.exists() and args.resume:
                write_json(contract_path, contracts[slide_number - 1])
                source_entry["slides"].append(
                    {
                        "slide": slide_number,
                        "category": category,
                        "topic": topic,
                        "structure": structure,
                        "template": relative_or_absolute(output, repo),
                        "constraints": relative_or_absolute(contract_path, repo),
                    }
                )
                print(
                    f"{relative_or_absolute(source, repo)}#{slide_number} -> "
                    f"{relative_or_absolute(output, repo)} (kept)"
                )
                continue
            if output.exists() and not args.overwrite:
                fail(f"Output already exists; pass --overwrite to replace it: {output}")
            with powerpoint_app() as extraction_app:
                extract_one(extraction_app, source, slide_number, output)
            write_json(contract_path, contracts[slide_number - 1])
            source_entry["slides"].append(
                {
                    "slide": slide_number,
                    "category": category,
                    "topic": topic,
                    "structure": structure,
                    "template": relative_or_absolute(output, repo),
                    "constraints": relative_or_absolute(
                        contract_path, repo
                    ),
                }
            )
            print(
                f"{relative_or_absolute(source, repo)}#{slide_number} -> "
                f"{relative_or_absolute(output, repo)}"
            )
        ledger["sources"].append(source_entry)
    candidate_count = sum(item["slide_count"] for item in ledger["sources"])
    ledger["validation"] = {
        "curation_status": "pending",
        "source_pages_reviewed": 0,
        "extracted_candidates": candidate_count,
        "curated_templates": 0,
        "merged_candidates": 0,
        "discarded_candidates": 0,
    }
    write_json(repo / "模板" / "学习记录.json", ledger)
    print(f"Wrote 模板/学习记录.json for {candidate_count} candidate slide(s)")
    print(
        "Curation is pending: review every page as keep/merge/discard and run "
        "script/curate_template_library.py --apply before indexing or archiving."
    )


def cmd_archive_learned(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    ledger_path = repo / "模板" / "学习记录.json"
    ledger = read_json(ledger_path)
    sources = ledger.get("sources")
    if not isinstance(sources, list) or not sources:
        fail("模板/学习记录.json contains no learned sources")
    if ledger.get("validation", {}).get("curation_status") != "passed":
        fail(
            "Cannot archive learned decks before keep/merge/discard curation passes"
        )
    inbox = (repo / "待提取").resolve()
    archive = (repo / "已提取").resolve()
    archive.mkdir(parents=True, exist_ok=True)
    moved = 0
    for source_entry in sources:
        source = resolve_repo_path(source_entry["source"], repo)
        if source.is_relative_to(archive):
            if not source.is_file():
                fail(f"Archived source is missing: {source}")
            destination = source
        else:
            if not source.is_relative_to(inbox):
                fail(f"Learned source is outside 待提取/ and 已提取/: {source}")
            if not source.is_file():
                fail(f"Learned source is missing: {source}")
            destination = archive / source.name
            if destination.exists():
                if sha256(destination) != sha256(source):
                    fail(f"Archive destination already contains a different file: {destination}")
                source.unlink()
            else:
                shutil.move(str(source), str(destination))
            moved += 1
        archived_path = relative_or_absolute(destination, repo)
        source_entry["source"] = archived_path
        source_entry["sha256"] = sha256(destination)
        for slide in source_entry.get("slides", []):
            if slide.get("curation", {}).get("decision") != "keep":
                continue
            contract_path = resolve_repo_path(slide["constraints"], repo)
            contract = read_json(contract_path)
            contract.setdefault("source", {})["deck"] = archived_path
            contract["source"]["sha256"] = source_entry["sha256"]
            write_json(contract_path, contract)
    write_json(ledger_path, ledger)
    print(f"Archived {moved} learned source deck(s) into 已提取/")


def load_plan(path: Path, repo: Path) -> list[dict[str, Any]]:
    data = read_json(path)
    slides = data.get("slides")
    if not isinstance(slides, list) or not slides:
        fail("Build plan must contain a non-empty 'slides' list")
    normalized: list[dict[str, Any]] = []
    dimensions: tuple[int, int] | None = None
    library = (repo / "模板").resolve()
    for index, entry in enumerate(slides, 1):
        if not isinstance(entry, dict) or not entry.get("template"):
            fail(f"Plan slide {index} is missing 'template'")
        template = resolve_repo_path(entry["template"], repo)
        if not template.is_relative_to(library):
            fail(f"Plan slide {index} is not sourced from 模板/: {template}")
        prs = open_presentation(template)
        if len(prs.slides) != 1:
            fail(
                f"Plan slide {index} uses a {len(prs.slides)}-slide file; "
                "curated templates must contain exactly one slide"
            )
        source_slide = int(entry.get("template_slide", 1))
        if source_slide < 1 or source_slide > len(prs.slides):
            fail(f"Plan slide {index} references invalid source slide {source_slide}")
        current_dimensions = (int(prs.slide_width), int(prs.slide_height))
        if dimensions is None:
            dimensions = current_dimensions
        elif current_dimensions != dimensions:
            fail(
                f"Mixed slide sizes: {template} has {current_dimensions}, expected {dimensions}"
            )
        normalized.append(
            {
                "template": template,
                "template_slide": source_slide,
                "purpose": str(entry.get("purpose", "")),
                "dimensions": current_dimensions,
            }
        )
    return normalized


def cmd_compose(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    plan_path = resolve_repo_path(args.plan, repo)
    output = resolve_repo_path(args.output, repo)
    if output.exists() and not args.overwrite:
        fail(f"Output already exists; pass --overwrite to replace it: {output}")
    plan = load_plan(plan_path, repo)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ppt-compose-") as temporary:
        temporary_seed = Path(temporary) / "seed.pptx"
        temporary_output = Path(temporary) / "deck.pptx"
        shutil.copy2(plan[0]["template"], temporary_seed)
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(
                str(temporary_seed), False, False, False
            )
            try:
                if presentation.Slides.Count != 1:
                    fail("The seed template must contain exactly one slide")
                for entry in plan[1:]:
                    source_slide = entry["template_slide"]
                    presentation.Slides.InsertFromFile(
                        str(entry["template"]),
                        presentation.Slides.Count,
                        source_slide,
                        source_slide,
                    )
                if presentation.Slides.Count != len(plan):
                    fail(
                        f"Composition produced {presentation.Slides.Count} slides; expected {len(plan)}"
                    )
                presentation.SaveAs(str(temporary_output), PPTX_FORMAT)
            finally:
                try:
                    presentation.Close()
                except Exception:
                    if not temporary_output.is_file():
                        raise
        copy_with_retry(temporary_output, output)
    provenance = {
        "schema_version": 1,
        "output": relative_or_absolute(output, repo),
        "build_plan": relative_or_absolute(plan_path, repo),
        "slides": [
            {
                "output_slide": number,
                "template": relative_or_absolute(entry["template"], repo),
                "template_slide": entry["template_slide"],
                "template_sha256": sha256(entry["template"]),
                "purpose": entry["purpose"],
            }
            for number, entry in enumerate(plan, 1)
        ],
    }
    sidecar = output.with_suffix(output.suffix + ".provenance.json")
    write_json(sidecar, provenance)
    print(output)
    print(sidecar)


def load_master_family(repo: Path, family_id: str) -> tuple[Path, dict[str, Any]]:
    index_path = repo / "模板" / "母版" / "索引.json"
    index = read_json(index_path)
    matches = [
        entry
        for entry in index.get("families", [])
        if isinstance(entry, dict) and str(entry.get("family_id")) == family_id
    ]
    if len(matches) != 1:
        fail(f"Expected one active master family '{family_id}' in {index_path}")
    entry = matches[0]
    if entry.get("status") != "active":
        fail(f"Master family '{family_id}' is not active")
    family_path = resolve_repo_path(entry["contract"], repo)
    family = read_json(family_path)
    if family.get("family_id") != family_id:
        fail(f"Family contract ID differs from index: {family_path}")
    return family_path, family


def replaceable_family_slots(
    role_data: dict[str, Any], roles: dict[str, Any] | None = None
) -> list[dict[str, Any]]:
    inherited_role = role_data.get("constraints_from_role")
    if inherited_role is not None:
        if roles is None or inherited_role not in roles:
            fail(
                f"Family role inherits constraints from unknown role '{inherited_role}'"
            )
        inherited_data = roles[inherited_role]
        if not isinstance(inherited_data, dict):
            fail(
                f"Family role constraint source '{inherited_role}' must be an object"
            )
        role_data = inherited_data
    return [
        copy.deepcopy(slot)
        for slot in role_data.get("text_slots", [])
        if isinstance(slot, dict) and "fixed_text" not in slot
    ]


def com_slide_by_id(presentation: Any, slide_id: int) -> Any:
    for index in range(1, int(presentation.Slides.Count) + 1):
        slide = presentation.Slides(index)
        if int(slide.SlideID) == slide_id:
            return slide
    fail(f"PowerPoint slide ID not found during family composition: {slide_id}")


def com_shape_refs_in_z_order(
    slide: Any, shape_refs: list[dict[str, Any]]
) -> list[Any]:
    requested = {
        (str(ref.get("name", "")), int(ref.get("occurrence", 1)))
        for ref in shape_refs
        if isinstance(ref, dict) and ref.get("name")
    }
    if len(requested) != len(shape_refs):
        fail("Module shape_refs contains an invalid or duplicate selector")
    occurrences: dict[str, int] = {}
    selected: list[Any] = []
    found: set[tuple[str, int]] = set()
    for index in range(1, int(slide.Shapes.Count) + 1):
        shape = slide.Shapes(index)
        name = str(shape.Name)
        occurrences[name] = occurrences.get(name, 0) + 1
        key = (name, occurrences[name])
        if key in requested:
            found.add(key)
            selected.append(shape)
    missing = sorted(requested - found)
    if missing:
        fail(f"Module source slide is missing registered shapes: {missing}")
    return selected


def mount_external_module(
    app: Any,
    target_slide: Any,
    module: dict[str, Any],
    module_source: Path,
    fixed_shell_names: set[str],
) -> None:
    source_presentation = app.Presentations.Open(
        str(module_source), True, False, False
    )
    try:
        mount_external_module_from_presentation(
            target_slide, module, source_presentation, fixed_shell_names
        )
    finally:
        source_presentation.Close()


def mount_external_module_from_presentation(
    target_slide: Any,
    module: dict[str, Any],
    source_presentation: Any,
    fixed_shell_names: set[str],
) -> None:
    """Mount a module using an already-open, read-only source presentation."""
    for index in range(int(target_slide.Shapes.Count), 0, -1):
        shape = target_slide.Shapes(index)
        if str(shape.Name) not in fixed_shell_names:
            shape.Delete()
    source_slide_number = int(module["source"]["slide"])
    if source_slide_number < 1 or source_slide_number > int(
        source_presentation.Slides.Count
    ):
        fail("External module source slide is outside its source deck")
    source_slide = source_presentation.Slides(source_slide_number)
    for source_shape in com_shape_refs_in_z_order(
        source_slide, module["shape_refs"]
    ):
        source_geometry = (
            float(source_shape.Left),
            float(source_shape.Top),
            float(source_shape.Width),
            float(source_shape.Height),
            float(source_shape.Rotation),
        )
        source_name = str(source_shape.Name)
        source_shape.Copy()
        pasted_range = target_slide.Shapes.Paste()
        pasted_shape = pasted_range.Item(1)
        pasted_geometry = (
            float(pasted_shape.Left),
            float(pasted_shape.Top),
            float(pasted_shape.Width),
            float(pasted_shape.Height),
            float(pasted_shape.Rotation),
        )
        geometry_drifted = any(
            abs(before - after) > 0.01
            for before, after in zip(source_geometry, pasted_geometry)
        )
        if geometry_drifted:
            # PowerPoint occasionally applies its standard 10-point repeated-paste
            # offset to visually identical shapes. Restoring the source values is
            # a mechanical normalization, not a relayout: the mounted contract
            # still requires the exact original absolute geometry.
            pasted_shape.Left = source_geometry[0]
            pasted_shape.Top = source_geometry[1]
            pasted_shape.Width = source_geometry[2]
            pasted_shape.Height = source_geometry[3]
            pasted_shape.Rotation = source_geometry[4]
            corrected_geometry = (
                float(pasted_shape.Left),
                float(pasted_shape.Top),
                float(pasted_shape.Width),
                float(pasted_shape.Height),
                float(pasted_shape.Rotation),
            )
            if any(
                abs(before - after) > 0.01
                for before, after in zip(source_geometry, corrected_geometry)
            ):
                fail(
                    f"PowerPoint changed geometry while mounting external shape '{source_name}': "
                    f"source={source_geometry}, pasted={pasted_geometry}, "
                    f"corrected={corrected_geometry}"
                )
        try:
            pasted_shape.Name = source_name
        except Exception as exc:
            fail(
                f"PowerPoint could not preserve mounted shape name '{source_name}': {exc}"
            )


def cmd_compose_family(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    plan_path = resolve_repo_path(args.plan, repo)
    plan = read_json(plan_path)
    family_id = str(plan.get("master_family", ""))
    if not family_id:
        fail("Layered build plan is missing master_family")
    family_path, family = load_master_family(repo, family_id)
    source_asset = resolve_repo_path(family["source_asset"], repo)
    if sha256(source_asset).upper() != str(family["source_sha256"]).upper():
        fail("Master-family source asset hash differs from family.json")
    source_deck = open_presentation(source_asset)
    slides = plan.get("slides")
    if not isinstance(slides, list) or not slides:
        fail("Layered build plan must contain a non-empty slides list")
    roles = family.get("role_pages", {})
    normalized: list[dict[str, Any]] = []
    for output_number, entry in enumerate(slides, 1):
        if not isinstance(entry, dict):
            fail(f"Plan slide {output_number} must be an object")
        shell = entry.get("shell", {})
        role = str(shell.get("role", ""))
        if role not in roles:
            fail(f"Plan slide {output_number} references unknown family role '{role}'")
        role_data = roles[role]
        shell_source_slide = int(
            shell.get("source_slide", role_data.get("source_slide", 0))
        )
        if shell_source_slide < 1 or shell_source_slide > len(source_deck.slides):
            fail(f"Plan slide {output_number} has invalid shell source slide")
        content = entry.get("content", {})
        content_type = str(content.get("type", ""))
        module_path: Path | None = None
        module: dict[str, Any] | None = None
        module_source: Path | None = None
        assembly = "same_family_source_slide_copy"
        source_slide = shell_source_slide
        capacity_slots = replaceable_family_slots(role_data, roles)
        if content_type == "family_role":
            content_role = str(content.get("role", role))
            if content_role != role:
                fail(
                    f"Plan slide {output_number} family_role '{content_role}' differs from shell role '{role}'"
                )
        elif content_type == "module":
            module_path = resolve_repo_path(content.get("module", ""), repo)
            module = read_json(module_path)
            if family_id not in module.get("compatible_master_families", []):
                fail(
                    f"Plan slide {output_number} module is not compatible with {family_id}"
                )
            module_source = resolve_repo_path(module["source"]["deck"], repo)
            if sha256(module_source).upper() != str(module["source"]["sha256"]).upper():
                fail(f"Plan slide {output_number} module source hash differs")
            module_deck = open_presentation(module_source)
            module_source_slide = int(module["source"]["slide"])
            if module_source_slide < 1 or module_source_slide > len(module_deck.slides):
                fail(f"Plan slide {output_number} module has invalid source slide")
            if (
                int(module_deck.slide_width) != int(source_deck.slide_width)
                or int(module_deck.slide_height) != int(source_deck.slide_height)
            ):
                fail(f"Plan slide {output_number} module page size differs from family")
            if module_source == source_asset:
                source_slide = module_source_slide
                source_layout = str(
                    source_deck.slides[source_slide - 1].slide_layout.name
                )
                if source_layout != str(role_data.get("layout", "")):
                    fail(
                        f"Plan slide {output_number} module source layout '{source_layout}' "
                        f"differs from shell layout '{role_data.get('layout')}'"
                    )
            else:
                source_slide = shell_source_slide
                assembly = "cross_source_module_mount"
            capacity_slots.extend(
                copy.deepcopy(module.get("constraints", {}).get("text_slots", []))
            )
        else:
            fail(
                f"Plan slide {output_number} content.type must be family_role or module"
            )
        normalized.append(
            {
                "purpose": str(entry.get("purpose", "")),
                "role": role,
                "shell_source_slide": shell_source_slide,
                "source_slide": source_slide,
                "assembly": assembly,
                "module_path": module_path,
                "module": module,
                "module_source": module_source,
                "capacity_contract": {
                    "schema_version": 1,
                    "text_slots": capacity_slots,
                },
            }
        )

    output = resolve_repo_path(args.output, repo)
    if output.exists() and not args.overwrite:
        fail(f"Output already exists; pass --overwrite to replace it: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    compose_state_path = output.with_suffix(output.suffix + ".compose-state.json")
    batch_path = output.with_suffix(output.suffix + ".compose-batch.pptx")
    plan_hash = sha256(plan_path).upper()
    state: dict[str, Any] | None = None
    if output.is_file() and compose_state_path.is_file():
        candidate = read_json(compose_state_path)
        if (
            str(candidate.get("plan_sha256", "")).upper() == plan_hash
            and str(candidate.get("master_source_sha256", "")).upper()
            == str(family["source_sha256"]).upper()
        ):
            state = candidate
    if state is None:
        shutil.copy2(source_asset, output)
        state = {
            "schema_version": 1,
            "plan_sha256": plan_hash,
            "master_source_sha256": str(family["source_sha256"]).upper(),
            "source_slide_ids": {},
            "shells_done": 0,
            "originals_deleted": False,
            "mounts_done": 0,
        }
        write_json(compose_state_path, state)

    if not state.get("source_slide_ids"):
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(
                str(output), True, False, False
            )
            try:
                source_slide_ids = {
                    number: int(presentation.Slides(number).SlideID)
                    for number in range(1, int(presentation.Slides.Count) + 1)
                }
            finally:
                try:
                    presentation.Close()
                except Exception:
                    pass
        state["source_slide_ids"] = {
            str(number): slide_id for number, slide_id in source_slide_ids.items()
        }
        write_json(compose_state_path, state)
    else:
        source_slide_ids = {
            int(number): int(slide_id)
            for number, slide_id in state["source_slide_ids"].items()
        }

    # Each batch is built from the last committed checkpoint.  A PowerPoint/RPC
    # failure can therefore be resumed without trusting a partially saved file.
    shell_batch_size = 3
    while int(state.get("shells_done", 0)) < len(normalized):
        start = int(state.get("shells_done", 0))
        batch = normalized[start : start + shell_batch_size]
        shutil.copy2(output, batch_path)
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(
                str(batch_path), False, False, False
            )
            try:
                for entry in batch:
                    original = com_slide_by_id(
                        presentation, source_slide_ids[entry["source_slide"]]
                    )
                    original_index = int(original.SlideIndex)
                    original.Duplicate()
                    presentation.Slides(original_index + 1).MoveTo(
                        presentation.Slides.Count
                    )
                presentation.Save()
            finally:
                try:
                    presentation.Close()
                except Exception:
                    pass
        copy_with_retry(batch_path, output)
        state["shells_done"] = start + len(batch)
        write_json(compose_state_path, state)

    if not bool(state.get("originals_deleted")):
        shutil.copy2(output, batch_path)
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(
                str(batch_path), False, False, False
            )
            try:
                for slide_id in source_slide_ids.values():
                    com_slide_by_id(presentation, slide_id).Delete()
                if int(presentation.Slides.Count) != len(normalized):
                    fail(
                        f"Family composition produced {presentation.Slides.Count} slides; "
                        f"expected {len(normalized)}"
                    )
                presentation.Save()
            finally:
                try:
                    presentation.Close()
                except Exception:
                    pass
        copy_with_retry(batch_path, output)
        state["originals_deleted"] = True
        write_json(compose_state_path, state)

    fixed_shell_names = {
        str(name)
        for name in family.get("shell_contract", {}).get(
            "fixed_slide_shapes", []
        )
    }
    mounts = [
        (slide_number, entry)
        for slide_number, entry in enumerate(normalized, 1)
        if entry["assembly"] == "cross_source_module_mount"
    ]
    mount_batch_size = 1
    while int(state.get("mounts_done", 0)) < len(mounts):
        start = int(state.get("mounts_done", 0))
        batch = mounts[start : start + mount_batch_size]
        shutil.copy2(output, batch_path)
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(
                str(batch_path), False, False, False
            )
            source_presentations: dict[Path, Any] = {}
            try:
                for slide_number, entry in batch:
                    if entry["module"] is None or entry["module_source"] is None:
                        fail("Cross-source module entry is incomplete")
                    module_source = entry["module_source"]
                    source_presentation = source_presentations.get(module_source)
                    if source_presentation is None:
                        source_presentation = app.Presentations.Open(
                            str(module_source), True, False, False
                        )
                        source_presentations[module_source] = source_presentation
                    mount_external_module_from_presentation(
                        presentation.Slides(slide_number),
                        entry["module"],
                        source_presentation,
                        fixed_shell_names,
                    )
                presentation.Save()
            finally:
                for source_presentation in source_presentations.values():
                    try:
                        source_presentation.Close()
                    except Exception:
                        pass
                try:
                    presentation.Close()
                except Exception:
                    pass
        copy_with_retry(batch_path, output)
        state["mounts_done"] = start + len(batch)
        write_json(compose_state_path, state)

    if batch_path.exists():
        batch_path.unlink()
    if compose_state_path.exists():
        compose_state_path.unlink()

    provenance = {
        "schema_version": 2,
        "output": relative_or_absolute(output, repo),
        "build_plan": relative_or_absolute(plan_path, repo),
        "master_family": family_id,
        "family_contract": relative_or_absolute(family_path, repo),
        "master_source": relative_or_absolute(source_asset, repo),
        "master_source_sha256": sha256(source_asset),
        "slides": [],
    }
    for output_number, entry in enumerate(normalized, 1):
        item: dict[str, Any] = {
            "output_slide": output_number,
            "purpose": entry["purpose"],
            "source_slide": entry["source_slide"],
            "assembly": entry["assembly"],
            "shell": {
                "role": entry["role"],
                "source_slide": entry["shell_source_slide"],
            },
            "capacity_contract": entry["capacity_contract"],
        }
        if entry["module_path"] is not None and entry["module"] is not None:
            item["content_module"] = {
                "contract": relative_or_absolute(entry["module_path"], repo),
                "source_slide": entry["module"]["source"]["slide"],
                "source_deck": entry["module"]["source"]["deck"],
                "source_sha256": entry["module"]["source"]["sha256"],
                "geometry_sha256": entry["module"]["geometry_sha256"],
            }
        else:
            item["content_role"] = entry["role"]
        provenance["slides"].append(item)
    sidecar = output.with_suffix(output.suffix + ".provenance.json")
    write_json(sidecar, provenance)
    print(output)
    print(sidecar)


def slide_xml_name(slide_number: int) -> str:
    return f"ppt/slides/slide{slide_number}.xml"


def slide_rels_name(slide_number: int) -> str:
    return f"ppt/slides/_rels/slide{slide_number}.xml.rels"


def parent_map(root: ET.Element) -> dict[ET.Element, ET.Element]:
    return {child: parent for parent in root.iter() for child in parent}


def named_shape_container(
    root: ET.Element, shape_name: str, occurrence: int | None = None
) -> ET.Element:
    parents = parent_map(root)
    candidates = [
        node
        for node in root.findall(".//p:cNvPr", NS)
        if node.attrib.get("name") == shape_name
    ]
    if occurrence is None:
        if len(candidates) != 1:
            fail(
                f"Expected one shape named '{shape_name}', found {len(candidates)}; "
                "set a 1-based 'occurrence' in the replacement map"
            )
        occurrence = 1
    if occurrence < 1 or occurrence > len(candidates):
        fail(
            f"Shape '{shape_name}' occurrence {occurrence} is outside "
            f"1..{len(candidates)}"
        )
    cursor = candidates[occurrence - 1]
    while cursor in parents:
        cursor = parents[cursor]
        if cursor.tag in {
            f"{{{NS['p']}}}sp",
            f"{{{NS['p']}}}pic",
            f"{{{NS['p']}}}graphicFrame",
        }:
            return cursor
    fail(f"Could not locate the XML container for shape '{shape_name}'")


def visible_char_count(value: str) -> int:
    return len(re.sub(r"\s", "", value))


def replacement_lines(entry: dict[str, Any]) -> tuple[list[str] | None, str]:
    if "rich_lines" in entry:
        rich_lines = entry["rich_lines"]
        if not isinstance(rich_lines, list) or not rich_lines:
            fail("Replacement 'rich_lines' must be a non-empty list")
        lines: list[str] = []
        for line_number, segments in enumerate(rich_lines, 1):
            if not isinstance(segments, list) or not segments:
                fail(f"Replacement rich line {line_number} must contain segments")
            parts: list[str] = []
            for segment_number, segment in enumerate(segments, 1):
                if not isinstance(segment, dict):
                    fail(
                        f"Replacement rich line {line_number} segment {segment_number} must be an object"
                    )
                text = segment.get("text")
                style = segment.get("style")
                if not isinstance(text, str) or not text:
                    fail(
                        f"Replacement rich line {line_number} segment {segment_number} needs non-empty text"
                    )
                if not isinstance(style, str) or not style:
                    fail(
                        f"Replacement rich line {line_number} segment {segment_number} needs a style role"
                    )
                parts.append(text)
            lines.append("".join(parts))
        return lines, "".join(lines)
    if "lines" in entry:
        lines = entry["lines"]
        if not isinstance(lines, list) or not lines or not all(
            isinstance(line, str) for line in lines
        ):
            fail("Replacement 'lines' must be a non-empty list of strings")
        return lines, "".join(lines)
    return None, str(entry.get("text", ""))


def validate_text_mapping(
    deck_path: Path, mapping: list[dict[str, Any]], repo: Path
) -> None:
    provenance_path = deck_path.with_suffix(deck_path.suffix + ".provenance.json")
    provenance = read_json(provenance_path)
    sources = provenance.get("slides")
    if not isinstance(sources, list) or not sources:
        fail("Text replacement requires a provenance sidecar with slide sources")
    by_slide: dict[int, dict[tuple[str, int | None], dict[str, Any]]] = {}
    for slide_number, source in enumerate(sources, 1):
        if int(provenance.get("schema_version", 1)) == 2:
            contract = source.get("capacity_contract")
            contract_path = provenance_path
            if not isinstance(contract, dict):
                fail(
                    f"Layered provenance slide {slide_number} has no capacity_contract"
                )
        else:
            template_path = resolve_repo_path(source["template"], repo)
            contract_path = constraints_path(template_path)
            if not contract_path.is_file():
                fail(f"Missing capacity contract for slide {slide_number}: {contract_path}")
            contract = read_json(contract_path)
        contract_errors = validate_constraints_shape(contract, contract_path)
        if contract_errors:
            fail("; ".join(contract_errors))
        slots: dict[tuple[str, int | None], dict[str, Any]] = {}
        for slot in contract["text_slots"]:
            occurrence = int(slot["occurrence"]) if "occurrence" in slot else None
            slots[(str(slot["shape"]), occurrence)] = slot
        by_slide[slide_number] = slots

    supplied: dict[int, set[tuple[str, int | None]]] = {
        number: set() for number in by_slide
    }
    for entry_number, entry in enumerate(mapping, 1):
        if not isinstance(entry, dict) or "slide" not in entry or "shape" not in entry:
            fail(f"Text replacement {entry_number} is missing slide or shape")
        slide_number = int(entry["slide"])
        if slide_number not in by_slide:
            fail(f"Text replacement {entry_number} references unknown slide {slide_number}")
        occurrence = int(entry["occurrence"]) if "occurrence" in entry else None
        key = (str(entry["shape"]), occurrence)
        if key in supplied[slide_number]:
            fail(f"Duplicate replacement for slide {slide_number}, selector {key}")
        supplied[slide_number].add(key)
        slot = by_slide[slide_number].get(key)
        if slot is None:
            fail(
                f"Slide {slide_number} replacement selector {key} is absent from its capacity contract"
            )
        lines, combined = replacement_lines(entry)
        style_roles = slot.get("style_roles")
        if style_roles is not None:
            rich_lines = entry.get("rich_lines")
            if not isinstance(rich_lines, list):
                fail(
                    f"Slide {slide_number} selector {key} preserves template emphasis; use 'rich_lines'"
                )
            used_roles = {
                str(segment["style"])
                for rich_line in rich_lines
                for segment in rich_line
            }
            unknown_roles = sorted(used_roles - set(style_roles))
            if unknown_roles:
                fail(
                    f"Slide {slide_number} selector {key} uses unknown style roles {unknown_roles}"
                )
            missing_roles = sorted(
                set(slot.get("required_style_roles", [])) - used_roles
            )
            if missing_roles:
                fail(
                    f"Slide {slide_number} selector {key} is missing required style roles {missing_roles}"
                )
        count = visible_char_count(combined)
        minimum = int(slot["min_chars"])
        maximum = int(slot["max_chars"])
        if count < minimum or count > maximum:
            fail(
                f"Slide {slide_number} selector {key} has {count} visible characters; "
                f"allowed range is {minimum}..{maximum}"
            )
        if "line_count" in slot:
            expected_lines = int(slot["line_count"])
            if lines is None:
                fail(
                    f"Slide {slide_number} selector {key} has a fixed line count; use 'lines' or 'rich_lines'"
                )
            if len(lines) != expected_lines:
                fail(
                    f"Slide {slide_number} selector {key} requires {expected_lines} line(s), "
                    f"received {len(lines)}"
                )
            line_mins = slot.get("line_min_chars", [0] * expected_lines)
            line_maxes = slot.get("line_max_chars", [maximum] * expected_lines)
            for line_number, line in enumerate(lines, 1):
                line_count = visible_char_count(line)
                if line_count < int(line_mins[line_number - 1]) or line_count > int(
                    line_maxes[line_number - 1]
                ):
                    fail(
                        f"Slide {slide_number} selector {key} line {line_number} has "
                        f"{line_count} visible characters; allowed range is "
                        f"{line_mins[line_number - 1]}..{line_maxes[line_number - 1]}"
                    )

    for slide_number, slots in by_slide.items():
        missing = [
            key
            for key, slot in slots.items()
            if slot.get("required", False) and key not in supplied[slide_number]
        ]
        if missing:
            fail(f"Slide {slide_number} is missing required text replacements: {missing}")


def replace_container_with_lines(
    container: ET.Element, lines: list[str], context: str = "target shape"
) -> None:
    paragraphs = container.findall(".//a:p", NS)
    if len(paragraphs) == len(lines) and len(paragraphs) > 1:
        for paragraph, line in zip(paragraphs, lines):
            text_nodes = paragraph.findall(".//a:t", NS)
            if not text_nodes:
                fail(f"{context}: paragraph has no replaceable text run")
            for text_number, node in enumerate(text_nodes):
                node.text = line if text_number == 0 else ""
        return
    if len(paragraphs) != 1:
        fail(
            f"{context}: fixed-line replacement requires one paragraph or exactly "
            f"one existing paragraph per line; found {len(paragraphs)} paragraph(s) "
            f"for {len(lines)} line(s)"
        )
    paragraph = paragraphs[0]
    run_tag = f"{{{NS['a']}}}r"
    field_tag = f"{{{NS['a']}}}fld"
    break_tag = f"{{{NS['a']}}}br"
    text_children = [
        child
        for child in list(paragraph)
        if child.tag in {run_tag, field_tag} and child.find(".//a:t", NS) is not None
    ]
    if not text_children:
        fail(f"{context}: fixed-line target has no text run to preserve")
    prototype = copy.deepcopy(text_children[0])
    insert_at = min(list(paragraph).index(child) for child in text_children)
    for child in list(paragraph):
        if child.tag in {run_tag, field_tag, break_tag}:
            paragraph.remove(child)
    for line_number, line in enumerate(lines):
        run = copy.deepcopy(prototype)
        text_nodes = run.findall(".//a:t", NS)
        for text_number, node in enumerate(text_nodes):
            node.text = line if text_number == 0 else ""
        paragraph.insert(insert_at, run)
        insert_at += 1
        if line_number < len(lines) - 1:
            paragraph.insert(insert_at, ET.Element(break_tag))
            insert_at += 1


def replace_container_with_rich_lines(
    container: ET.Element,
    prototype_container: ET.Element,
    rich_lines: list[list[dict[str, str]]],
    style_roles: dict[str, int],
    context: str = "target shape",
) -> None:
    paragraphs = container.findall(".//a:p", NS)
    if len(paragraphs) != 1:
        fail(
            f"{context}: rich-line replacement requires exactly one paragraph; "
            f"found {len(paragraphs)}"
        )
    paragraph = paragraphs[0]
    run_tag = f"{{{NS['a']}}}r"
    field_tag = f"{{{NS['a']}}}fld"
    break_tag = f"{{{NS['a']}}}br"
    text_children = [
        child
        for child in list(paragraph)
        if child.tag in {run_tag, field_tag} and child.find(".//a:t", NS) is not None
    ]
    if not text_children:
        fail("Rich-line target has no text runs to preserve")
    prototype_paragraphs = prototype_container.findall(".//a:p", NS)
    if len(prototype_paragraphs) != 1:
        fail("Rich-line style source requires exactly one paragraph")
    prototype_text_children = [
        child
        for child in list(prototype_paragraphs[0])
        if child.tag in {run_tag, field_tag} and child.find(".//a:t", NS) is not None
    ]
    if not prototype_text_children:
        fail("Rich-line style source has no text runs")
    prototypes: dict[str, ET.Element] = {}
    for role, source_run in style_roles.items():
        if source_run > len(prototype_text_children):
            fail(
                f"Style role '{role}' references source run {source_run}, "
                f"but the template source has only {len(prototype_text_children)} text runs"
            )
        prototypes[role] = copy.deepcopy(prototype_text_children[source_run - 1])
    insert_at = min(list(paragraph).index(child) for child in text_children)
    for child in list(paragraph):
        if child.tag in {run_tag, field_tag, break_tag}:
            paragraph.remove(child)
    for line_number, segments in enumerate(rich_lines):
        for segment in segments:
            role = str(segment["style"])
            if role not in prototypes:
                fail(f"Unknown style role '{role}'")
            run = copy.deepcopy(prototypes[role])
            text_nodes = run.findall(".//a:t", NS)
            for text_number, node in enumerate(text_nodes):
                node.text = str(segment["text"]) if text_number == 0 else ""
            paragraph.insert(insert_at, run)
            insert_at += 1
        if line_number < len(rich_lines) - 1:
            paragraph.insert(insert_at, ET.Element(break_tag))
            insert_at += 1


def pptx_patch(input_path: Path, output_path: Path, patcher: Any) -> None:
    if not input_path.is_file():
        fail(f"PPTX not found: {input_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if input_path.resolve() != output_path.resolve():
        shutil.copy2(input_path, output_path)
    temp_handle, temp_name = tempfile.mkstemp(
        prefix=output_path.stem + ".", suffix=".pptx", dir=output_path.parent
    )
    os.close(temp_handle)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(output_path, "r") as source_zip:
            names = set(source_zip.namelist())
            replacements, additions = patcher(source_zip, names)
            with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target_zip:
                for item in source_zip.infolist():
                    payload = replacements.get(item.filename, source_zip.read(item.filename))
                    target_zip.writestr(item, payload)
                for name, payload in additions.items():
                    target_zip.writestr(name, payload)
        last_error: PermissionError | None = None
        for attempt in range(12):
            try:
                os.replace(temp_path, output_path)
                last_error = None
                break
            except PermissionError as exc:
                last_error = exc
                time.sleep(0.25 * (attempt + 1))
        if last_error is not None:
            raise last_error
    finally:
        if temp_path.exists():
            temp_path.unlink()


def cmd_replace_text(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    input_path = resolve_repo_path(args.input, repo)
    output_path = resolve_repo_path(args.output, repo) if args.output else input_path
    mapping = read_json(resolve_repo_path(args.map, repo)).get("replacements")
    if not isinstance(mapping, list) or not mapping:
        fail("Text map must contain a non-empty 'replacements' list")
    validate_text_mapping(input_path, mapping, repo)
    provenance_path = input_path.with_suffix(input_path.suffix + ".provenance.json")
    provenance = read_json(provenance_path)

    def patcher(source_zip: zipfile.ZipFile, names: set[str]) -> tuple[dict[str, bytes], dict[str, bytes]]:
        replacements: dict[str, bytes] = {}
        by_slide: dict[int, list[dict[str, Any]]] = {}
        for entry in mapping:
            by_slide.setdefault(int(entry["slide"]), []).append(entry)
        for slide_number, entries in by_slide.items():
            xml_name = slide_xml_name(slide_number)
            if xml_name not in names:
                fail(f"Slide {slide_number} does not exist")
            root = ET.fromstring(source_zip.read(xml_name))
            for entry in entries:
                shape_name = str(entry["shape"])
                occurrence = (
                    int(entry["occurrence"]) if "occurrence" in entry else None
                )
                output_occurrence = (
                    int(entry["output_occurrence"])
                    if "output_occurrence" in entry
                    else occurrence
                )
                container = named_shape_container(root, shape_name, output_occurrence)
                lines, _ = replacement_lines(entry)
                if "rich_lines" in entry:
                    source = provenance["slides"][slide_number - 1]
                    if int(provenance.get("schema_version", 1)) == 2:
                        template_path = resolve_repo_path(
                            provenance["master_source"], repo
                        )
                        template_slide = int(source["source_slide"])
                        contract = source["capacity_contract"]
                        module_source = source.get("content_module")
                        if module_source:
                            module = read_json(
                                resolve_repo_path(module_source["contract"], repo)
                            )
                            module_contract = module.get("constraints", {})
                            selector = (shape_name, occurrence)
                            module_selectors = {
                                (
                                    str(slot.get("shape", "")),
                                    int(slot["occurrence"])
                                    if "occurrence" in slot
                                    else None,
                                )
                                for slot in module_contract.get("text_slots", [])
                                if isinstance(slot, dict)
                            }
                            if selector in module_selectors:
                                template_path = resolve_repo_path(
                                    module["source"]["deck"], repo
                                )
                                template_slide = int(module["source"]["slide"])
                                contract = module_contract
                    else:
                        template_path = resolve_repo_path(source["template"], repo)
                        template_slide = int(source.get("template_slide", 1))
                        contract = read_json(constraints_path(template_path))
                    selector = (shape_name, occurrence)
                    matching_slots = [
                        slot
                        for slot in contract["text_slots"]
                        if (
                            str(slot["shape"]),
                            int(slot["occurrence"]) if "occurrence" in slot else None,
                        )
                        == selector
                    ]
                    if len(matching_slots) != 1:
                        fail(
                            f"Could not resolve one rich-text capacity slot for slide {slide_number}, selector {selector}"
                        )
                    with zipfile.ZipFile(template_path, "r") as template_zip:
                        prototype_root = ET.fromstring(
                            template_zip.read(
                                slide_xml_name(template_slide)
                            )
                        )
                    prototype_container = named_shape_container(
                        prototype_root, shape_name, occurrence
                    )
                    replace_container_with_rich_lines(
                        container,
                        prototype_container,
                        entry["rich_lines"],
                        matching_slots[0]["style_roles"],
                        f"slide {slide_number} shape '{shape_name}'",
                    )
                    continue
                if lines is not None:
                    replace_container_with_lines(
                        container,
                        lines,
                        f"slide {slide_number} shape '{shape_name}'",
                    )
                    continue
                paragraphs = container.findall(".//a:p", NS)
                new_lines = str(entry.get("text", "")).replace("\r\n", "\n").split("\n")
                if len(new_lines) > len(paragraphs):
                    fail(
                        f"'{shape_name}' on slide {slide_number} has {len(paragraphs)} paragraph slot(s), "
                        f"but replacement needs {len(new_lines)}"
                    )
                for index, paragraph in enumerate(paragraphs):
                    text_nodes = paragraph.findall(".//a:t", NS)
                    if not text_nodes and index < len(new_lines):
                        fail(f"Paragraph {index + 1} in '{shape_name}' has no replaceable text run")
                    value = new_lines[index] if index < len(new_lines) else ""
                    for text_index, node in enumerate(text_nodes):
                        node.text = value if text_index == 0 else ""
            replacements[xml_name] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )
        return replacements, {}

    pptx_patch(input_path, output_path, patcher)
    print(output_path)


def next_media_name(names: set[str], extension: str) -> str:
    numbers = []
    pattern = re.compile(r"ppt/media/image(\d+)\.[^.]+$", re.IGNORECASE)
    for name in names:
        match = pattern.match(name)
        if match:
            numbers.append(int(match.group(1)))
    return f"ppt/media/image{max(numbers, default=0) + 1}{extension}"


def next_relationship_id(rels_root: ET.Element) -> str:
    used = set()
    for rel in rels_root:
        match = re.fullmatch(r"rId(\d+)", rel.attrib.get("Id", ""))
        if match:
            used.add(int(match.group(1)))
    number = 1
    while number in used:
        number += 1
    return f"rId{number}"


def set_cover_crop(
    picture: ET.Element, image_path: Path, focal_x: float, focal_y: float
) -> None:
    xfrm = picture.find("./p:spPr/a:xfrm/a:ext", NS)
    if xfrm is None:
        fail("Picture shape has no fixed frame geometry")
    frame_width = int(xfrm.attrib["cx"])
    frame_height = int(xfrm.attrib["cy"])
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    if not image_width or not image_height or not frame_width or not frame_height:
        fail(f"Invalid image or frame dimensions for {image_path}")
    frame_ratio = frame_width / frame_height
    image_ratio = image_width / image_height
    left = top = right = bottom = 0.0
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        left = min(max(focal_x - visible / 2, 0.0), 1.0 - visible)
        right = 1.0 - visible - left
    elif image_ratio < frame_ratio:
        visible = image_ratio / frame_ratio
        top = min(max(focal_y - visible / 2, 0.0), 1.0 - visible)
        bottom = 1.0 - visible - top
    blip_fill = picture.find("./p:blipFill", NS)
    if blip_fill is None:
        fail("Picture shape has no blipFill")
    src_rect = blip_fill.find("./a:srcRect", NS)
    if src_rect is None:
        src_rect = ET.Element(f"{{{NS['a']}}}srcRect")
        blip = blip_fill.find("./a:blip", NS)
        insert_at = list(blip_fill).index(blip) + 1 if blip is not None else 0
        blip_fill.insert(insert_at, src_rect)
    src_rect.attrib.clear()
    for key, value in (("l", left), ("t", top), ("r", right), ("b", bottom)):
        amount = round(value * 100000)
        if amount:
            src_rect.set(key, str(amount))


def cmd_replace_image(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    input_path = resolve_repo_path(args.input, repo)
    output_path = resolve_repo_path(args.output, repo) if args.output else input_path
    mapping = read_json(resolve_repo_path(args.map, repo)).get("replacements")
    if not isinstance(mapping, list) or not mapping:
        fail("Image map must contain a non-empty 'replacements' list")

    def patcher(source_zip: zipfile.ZipFile, names: set[str]) -> tuple[dict[str, bytes], dict[str, bytes]]:
        replacements: dict[str, bytes] = {}
        additions: dict[str, bytes] = {}
        content_types = ET.fromstring(source_zip.read("[Content_Types].xml"))
        by_slide: dict[int, list[dict[str, Any]]] = {}
        for entry in mapping:
            by_slide.setdefault(int(entry["slide"]), []).append(entry)
        occupied_names = set(names)
        for slide_number, entries in by_slide.items():
            xml_name = slide_xml_name(slide_number)
            rels_name = slide_rels_name(slide_number)
            if xml_name not in names or rels_name not in names:
                fail(f"Slide {slide_number} or its relationships do not exist")
            slide_root = ET.fromstring(source_zip.read(xml_name))
            rels_root = ET.fromstring(source_zip.read(rels_name))
            for entry in entries:
                image_path = resolve_repo_path(entry["image"], repo)
                extension = image_path.suffix.lower()
                if extension not in CONTENT_TYPES:
                    fail(f"Unsupported image type '{extension}' for {image_path}")
                if not image_path.is_file():
                    fail(f"Replacement image not found: {image_path}")
                focal_x = float(entry.get("focal_x", 0.5))
                focal_y = float(entry.get("focal_y", 0.5))
                if not 0 <= focal_x <= 1 or not 0 <= focal_y <= 1:
                    fail("focal_x and focal_y must be between 0 and 1")
                shape_name = str(entry["shape"])
                occurrence = (
                    int(entry["occurrence"]) if "occurrence" in entry else None
                )
                picture = named_shape_container(
                    slide_root, shape_name, occurrence
                )
                if picture.tag != f"{{{NS['p']}}}pic":
                    fail(f"'{shape_name}' on slide {slide_number} is not a picture shape")
                blip = picture.find("./p:blipFill/a:blip", NS)
                if blip is None:
                    fail(f"'{shape_name}' has no embedded image")
                new_rid = next_relationship_id(rels_root)
                media_name = next_media_name(occupied_names, extension)
                occupied_names.add(media_name)
                relationship = ET.SubElement(
                    rels_root, f"{{{NS['rel']}}}Relationship"
                )
                relationship.set("Id", new_rid)
                relationship.set("Type", IMAGE_REL_TYPE)
                relationship.set("Target", "../media/" + Path(media_name).name)
                blip.set(f"{{{NS['r']}}}embed", new_rid)
                blip.attrib.pop(f"{{{NS['r']}}}link", None)
                extension_list = blip.find("./a:extLst", NS)
                if extension_list is not None:
                    blip.remove(extension_list)
                if entry.get("fit", "cover") == "cover":
                    set_cover_crop(picture, image_path, focal_x, focal_y)
                elif entry.get("fit") != "preserve":
                    fail("Image fit must be 'cover' or 'preserve'")
                additions[media_name] = image_path.read_bytes()
                matching_defaults = [
                    node
                    for node in content_types.findall("ct:Default", NS)
                    if node.attrib.get("Extension", "").lower()
                    == extension.lstrip(".")
                ]
                if matching_defaults:
                    for default in matching_defaults:
                        default.set("ContentType", CONTENT_TYPES[extension])
                else:
                    default = ET.SubElement(
                        content_types, f"{{{NS['ct']}}}Default"
                    )
                    default.set("Extension", extension.lstrip("."))
                    default.set("ContentType", CONTENT_TYPES[extension])
            replacements[xml_name] = ET.tostring(
                slide_root, encoding="utf-8", xml_declaration=True
            )
            replacements[rels_name] = ET.tostring(
                rels_root, encoding="utf-8", xml_declaration=True
            )
        replacements["[Content_Types].xml"] = ET.tostring(
            content_types, encoding="utf-8", xml_declaration=True
        )
        return replacements, additions

    pptx_patch(input_path, output_path, patcher)
    print(output_path)


def flatten_shape_signatures(shapes: list[dict[str, Any]], prefix: str = "") -> list[dict[str, Any]]:
    flattened: list[dict[str, Any]] = []
    for shape in shapes:
        path = f"{prefix}/{shape['name']}" if prefix else shape["name"]
        flattened.append(
            {
                "path": path,
                "type": shape["type"],
                "left": shape["left"],
                "top": shape["top"],
                "width": shape["width"],
                "height": shape["height"],
                "rotation": shape["rotation"],
            }
        )
        flattened.extend(flatten_shape_signatures(shape.get("children", []), path))
    return flattened


def cmd_verify(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    deck_path = resolve_repo_path(args.input, repo)
    sidecar = (
        resolve_repo_path(args.provenance, repo)
        if args.provenance
        else deck_path.with_suffix(deck_path.suffix + ".provenance.json")
    )
    provenance = read_json(sidecar)
    sources = provenance.get("slides")
    if not isinstance(sources, list) or not sources:
        fail("Provenance contains no slide sources")
    deck = open_presentation(deck_path)
    errors: list[str] = []
    if len(deck.slides) != len(sources):
        errors.append(f"output has {len(deck.slides)} slides; provenance has {len(sources)}")
    for index, source in enumerate(sources, 1):
        if index > len(deck.slides):
            break
        template_path = resolve_repo_path(source["template"], repo)
        if not template_path.is_relative_to((repo / "模板").resolve()):
            errors.append(f"slide {index}: source is outside 模板/")
            continue
        if sha256(template_path) != source.get("template_sha256"):
            errors.append(f"slide {index}: source template changed after composition")
        template = open_presentation(template_path)
        source_slide_number = int(source.get("template_slide", 1))
        if int(deck.slide_width) != int(template.slide_width) or int(deck.slide_height) != int(template.slide_height):
            errors.append(f"slide {index}: page size differs from source template")
        output_signature = flatten_shape_signatures(slide_record(deck.slides[index - 1], index)["shapes"])
        source_signature = flatten_shape_signatures(
            slide_record(template.slides[source_slide_number - 1], source_slide_number)["shapes"]
        )
        if output_signature != source_signature:
            errors.append(f"slide {index}: shape structure or geometry differs from source template")
    result = {
        "ok": not errors,
        "deck": relative_or_absolute(deck_path, repo),
        "slides_checked": min(len(deck.slides), len(sources)),
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(2)


def cmd_verify_layered(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    deck_path = resolve_repo_path(args.input, repo)
    sidecar = (
        resolve_repo_path(args.provenance, repo)
        if args.provenance
        else deck_path.with_suffix(deck_path.suffix + ".provenance.json")
    )
    provenance = read_json(sidecar)
    if int(provenance.get("schema_version", 0)) != 2:
        fail("Layered verification requires provenance schema_version 2")
    sources = provenance.get("slides")
    if not isinstance(sources, list) or not sources:
        fail("Layered provenance contains no slide sources")
    family_path = resolve_repo_path(provenance["family_contract"], repo)
    family = read_json(family_path)
    master_source = resolve_repo_path(provenance["master_source"], repo)
    errors: list[str] = []
    if str(family.get("family_id")) != str(provenance.get("master_family")):
        errors.append("provenance master_family differs from family contract")
    expected_master_sha = str(provenance.get("master_source_sha256", ""))
    actual_master_sha = sha256(master_source)
    if expected_master_sha.upper() != actual_master_sha.upper():
        errors.append("master source changed after family composition")
    deck = open_presentation(deck_path)
    source_deck = open_presentation(master_source)
    if len(deck.slides) != len(sources):
        errors.append(
            f"output has {len(deck.slides)} slides; provenance has {len(sources)}"
        )
    for index, source in enumerate(sources, 1):
        if index > len(deck.slides):
            break
        source_slide_number = int(source.get("source_slide", 0))
        if source_slide_number < 1 or source_slide_number > len(source_deck.slides):
            errors.append(f"slide {index}: invalid family source slide")
            continue
        assembly = str(source.get("assembly", "same_family_source_slide_copy"))
        module_source = source.get("content_module")
        if assembly == "cross_source_module_mount":
            if not module_source:
                errors.append(f"slide {index}: cross-source mount has no module provenance")
                continue
            shell_source_number = int(source.get("shell", {}).get("source_slide", 0))
            if shell_source_number < 1 or shell_source_number > len(source_deck.slides):
                errors.append(f"slide {index}: invalid shell source slide")
                continue
            fixed_refs = [
                {"name": str(name), "occurrence": 1}
                for name in family.get("shell_contract", {}).get(
                    "fixed_slide_shapes", []
                )
            ]
            shell_records, _, shell_selection_errors = module_geometry_records(
                source_deck.slides[shell_source_number - 1], fixed_refs
            )
            output_shell_records, _, output_shell_errors = module_geometry_records(
                deck.slides[index - 1], fixed_refs
            )
            if shell_selection_errors or output_shell_errors:
                errors.extend(
                    f"slide {index}: {message}"
                    for message in shell_selection_errors + output_shell_errors
                )
            elif mounted_geometry_records(shell_records) != mounted_geometry_records(
                output_shell_records
            ):
                errors.append(
                    f"slide {index}: target-master shell geometry differs from family source"
                )

            module_path = resolve_repo_path(module_source["contract"], repo)
            module = read_json(module_path)
            if (
                str(module.get("geometry_sha256", "")).upper()
                != str(module_source.get("geometry_sha256", "")).upper()
            ):
                errors.append(f"slide {index}: module geometry signature is stale")
            external_deck_path = resolve_repo_path(module["source"]["deck"], repo)
            if sha256(external_deck_path).upper() != str(
                module["source"].get("sha256", "")
            ).upper():
                errors.append(f"slide {index}: external module source changed")
                continue
            external_deck = open_presentation(external_deck_path)
            external_slide_number = int(module["source"].get("slide", 0))
            if external_slide_number < 1 or external_slide_number > len(
                external_deck.slides
            ):
                errors.append(f"slide {index}: external module source slide is invalid")
                continue
            shape_refs = module.get("shape_refs", [])
            source_module_records, _, source_module_errors = module_geometry_records(
                external_deck.slides[external_slide_number - 1], shape_refs
            )
            fixed_name_counts: dict[str, int] = {}
            for ref in fixed_refs:
                fixed_name = str(ref.get("name", ""))
                fixed_name_counts[fixed_name] = fixed_name_counts.get(fixed_name, 0) + 1
            output_shape_refs = [
                {
                    "name": str(ref.get("name", "")),
                    "occurrence": int(ref.get("occurrence", 1))
                    + fixed_name_counts.get(str(ref.get("name", "")), 0),
                }
                for ref in shape_refs
            ]
            output_module_records, _, output_module_errors = module_geometry_records(
                deck.slides[index - 1], output_shape_refs
            )
            normalized_output_module_records: list[dict[str, Any]] = []
            for record in output_module_records:
                normalized_record = copy.deepcopy(record)
                normalized_record["occurrence"] = int(record["occurrence"]) - fixed_name_counts.get(
                    str(record["name"]), 0
                )
                normalized_output_module_records.append(normalized_record)
            if source_module_errors or output_module_errors:
                errors.extend(
                    f"slide {index}: {message}"
                    for message in source_module_errors + output_module_errors
                )
            elif mounted_geometry_records(
                source_module_records
            ) != mounted_geometry_records(normalized_output_module_records):
                errors.append(
                    f"slide {index}: mounted module shape structure or geometry differs from its external source"
                )
            elif [record["z_index"] for record in output_module_records] != sorted(
                record["z_index"] for record in output_module_records
            ):
                errors.append(f"slide {index}: mounted module z-order is not preserved")
            expected_shape_count = len(fixed_refs) + len(shape_refs)
            if len(deck.slides[index - 1].shapes) != expected_shape_count:
                errors.append(
                    f"slide {index}: mounted slide has {len(deck.slides[index - 1].shapes)} "
                    f"top-level shapes; expected {expected_shape_count}"
                )
        else:
            output_signature = flatten_shape_signatures(
                slide_record(deck.slides[index - 1], index)["shapes"]
            )
            source_signature = flatten_shape_signatures(
                slide_record(
                    source_deck.slides[source_slide_number - 1], source_slide_number
                )["shapes"]
            )
            if output_signature != source_signature:
                errors.append(
                    f"slide {index}: shape structure or geometry differs from family source slide {source_slide_number}"
                )
            if module_source:
                module_path = resolve_repo_path(module_source["contract"], repo)
                module = read_json(module_path)
                if (
                    str(module.get("geometry_sha256", "")).upper()
                    != str(module_source.get("geometry_sha256", "")).upper()
                ):
                    errors.append(f"slide {index}: module geometry signature is stale")
                if int(module.get("source", {}).get("slide", 0)) != source_slide_number:
                    errors.append(f"slide {index}: module source slide differs from provenance")
    result = {
        "ok": not errors,
        "deck": relative_or_absolute(deck_path, repo),
        "master_family": provenance.get("master_family"),
        "slides_checked": min(len(deck.slides), len(sources)),
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(2)


def package_member_sha256(path: Path, member: str) -> str:
    normalized = member.lstrip("/")
    try:
        with zipfile.ZipFile(path, "r") as package:
            payload = package.read(normalized)
    except KeyError:
        fail(f"Package member not found in {path}: /{normalized}")
    return hashlib.sha256(payload).hexdigest()


def master_profile(deck_path: Path) -> dict[str, Any]:
    deck = open_presentation(deck_path)
    slides: list[dict[str, Any]] = []
    master_parts: set[str] = set()
    layout_usage: dict[str, int] = {}
    for number, slide in enumerate(deck.slides, 1):
        layout = slide.slide_layout
        master_part = str(layout.slide_master.part.partname)
        layout_part = str(layout.part.partname)
        layout_name = str(layout.name or "")
        master_parts.add(master_part)
        layout_usage[layout_name] = layout_usage.get(layout_name, 0) + 1
        slides.append(
            {
                "slide": number,
                "master_part": master_part,
                "layout_part": layout_part,
                "layout_name": layout_name,
            }
        )
    return {
        "deck": str(deck_path),
        "slide_count": len(deck.slides),
        "slide_size_emu": {
            "width": int(deck.slide_width),
            "height": int(deck.slide_height),
        },
        "master_count": len(master_parts),
        "master_parts": sorted(master_parts),
        "layout_usage": dict(sorted(layout_usage.items())),
        "slides": slides,
    }


def cmd_verify_master(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    deck_path = resolve_repo_path(args.input, repo)
    profile = master_profile(deck_path)
    errors: list[str] = []
    if profile["slide_count"] < 1:
        errors.append("deck contains no slides")
    if profile["master_count"] != 1:
        errors.append(
            f"deck uses {profile['master_count']} PowerPoint masters; expected exactly 1"
        )

    family_summary: dict[str, Any] | None = None
    if args.family:
        family_path = resolve_repo_path(args.family, repo)
        family = read_json(family_path)
        source_asset_value = family.get("source_asset")
        if not source_asset_value:
            fail(f"Family contract has no source_asset: {family_path}")
        source_asset = resolve_repo_path(source_asset_value, repo)
        source_profile = master_profile(source_asset)
        approved_layouts = {
            str(entry.get("name", ""))
            for entry in family.get("approved_layouts", [])
            if isinstance(entry, dict) and entry.get("name")
        }
        used_layouts = set(profile["layout_usage"])
        unapproved = sorted(used_layouts - approved_layouts)
        if unapproved:
            errors.append(f"unapproved slide layouts: {', '.join(unapproved)}")
        expected_size = family.get("slide_size_emu", {})
        actual_size = profile["slide_size_emu"]
        if (
            int(expected_size.get("width", -1)) != actual_size["width"]
            or int(expected_size.get("height", -1)) != actual_size["height"]
        ):
            errors.append(
                f"slide size {actual_size} differs from family {expected_size}"
            )
        if source_profile["master_count"] != 1:
            errors.append(
                f"family source uses {source_profile['master_count']} masters; expected 1"
            )
        elif profile["master_count"] == 1:
            deck_master = profile["master_parts"][0]
            source_master = source_profile["master_parts"][0]
            deck_master_sha = package_member_sha256(deck_path, deck_master)
            source_master_sha = package_member_sha256(source_asset, source_master)
            if deck_master_sha != source_master_sha:
                errors.append("output master definition differs from family source")
        expected_source_sha = str(family.get("source_sha256", ""))
        actual_source_sha = sha256(source_asset)
        if expected_source_sha and actual_source_sha.upper() != expected_source_sha.upper():
            errors.append("family source_asset hash differs from family.json")
        family_summary = {
            "family_id": family.get("family_id"),
            "contract": relative_or_absolute(family_path, repo),
            "source_asset": relative_or_absolute(source_asset, repo),
            "approved_layouts": sorted(approved_layouts),
        }

    result = {
        "ok": not errors,
        "deck": relative_or_absolute(deck_path, repo),
        "slide_count": profile["slide_count"],
        "master_count": profile["master_count"],
        "master_parts": profile["master_parts"],
        "layout_usage": profile["layout_usage"],
        "family": family_summary,
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(2)


def module_geometry_records(
    slide: Any, shape_refs: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[Any], list[str]]:
    requested = {
        (str(ref.get("name", "")), int(ref.get("occurrence", 1)))
        for ref in shape_refs
        if isinstance(ref, dict) and ref.get("name")
    }
    errors: list[str] = []
    if len(requested) != len(shape_refs):
        errors.append("shape_refs contains an invalid or duplicate shape reference")
    occurrences: dict[str, int] = {}
    records: list[dict[str, Any]] = []
    selected: list[Any] = []
    found: set[tuple[str, int]] = set()
    for z_index, shape in enumerate(slide.shapes, 1):
        name = str(shape.name)
        occurrences[name] = occurrences.get(name, 0) + 1
        key = (name, occurrences[name])
        if key not in requested:
            continue
        found.add(key)
        selected.append(shape)
        records.append(
            {
                "name": name,
                "occurrence": occurrences[name],
                "type": str(shape.shape_type),
                "left": int(shape.left),
                "top": int(shape.top),
                "width": int(shape.width),
                "height": int(shape.height),
                "rotation": float(shape.rotation or 0),
                "z_index": z_index,
            }
        )
    for name, occurrence in sorted(requested - found):
        errors.append(f"source slide has no top-level shape '{name}' occurrence {occurrence}")
    return records, selected, errors


def module_geometry_sha256(records: list[dict[str, Any]]) -> str:
    payload = json.dumps(
        records, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest().upper()


def mounted_geometry_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {key: value for key, value in record.items() if key != "z_index"}
        for record in records
    ]


def cmd_verify_module(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    module_path = resolve_repo_path(args.module, repo)
    module = read_json(module_path)
    errors: list[str] = []
    source = module.get("source", {})
    source_deck_value = source.get("deck")
    if not source_deck_value:
        fail(f"Module contract has no source.deck: {module_path}")
    source_deck = resolve_repo_path(source_deck_value, repo)
    if not source_deck.is_file():
        fail(f"Module source deck not found: {source_deck}")
    expected_source_sha = str(source.get("sha256", ""))
    actual_source_sha = sha256(source_deck)
    if expected_source_sha and actual_source_sha.upper() != expected_source_sha.upper():
        errors.append("module source deck hash differs from contract")
    expected_source_master_sha = str(source.get("master_sha256", ""))
    actual_source_master_sha = ""
    source_profile = master_profile(source_deck)
    if source_profile["master_count"] == 1:
        actual_source_master_sha = package_member_sha256(
            source_deck, source_profile["master_parts"][0]
        ).upper()
    elif expected_source_master_sha:
        errors.append(
            f"module source uses {source_profile['master_count']} masters; expected one fingerprinted master"
        )
    if (
        expected_source_master_sha
        and actual_source_master_sha != expected_source_master_sha.upper()
    ):
        errors.append("module source master fingerprint differs from contract")
    deck = open_presentation(source_deck)
    source_slide_number = int(source.get("slide", 0))
    if source_slide_number < 1 or source_slide_number > len(deck.slides):
        fail(
            f"Module source slide {source_slide_number} is outside 1..{len(deck.slides)}"
        )
    shape_refs = module.get("shape_refs")
    if not isinstance(shape_refs, list) or not shape_refs:
        fail("Module contract must contain a non-empty shape_refs list")
    records, selected, selection_errors = module_geometry_records(
        deck.slides[source_slide_number - 1], shape_refs
    )
    errors.extend(selection_errors)
    actual_geometry_sha = module_geometry_sha256(records)
    expected_geometry_sha = str(module.get("geometry_sha256", "")).upper()
    if expected_geometry_sha != actual_geometry_sha:
        errors.append("module geometry signature differs from source slide")
    if selected:
        actual_bounds = {
            "left": min(int(shape.left) for shape in selected),
            "top": min(int(shape.top) for shape in selected),
            "right": max(int(shape.left + shape.width) for shape in selected),
            "bottom": max(int(shape.top + shape.height) for shape in selected),
        }
    else:
        actual_bounds = None
    if actual_bounds != module.get("bounds_emu"):
        errors.append(
            f"module bounds {actual_bounds} differ from contract {module.get('bounds_emu')}"
        )

    constraints = module.get("constraints", {})
    text_slots = constraints.get("text_slots", [])
    if not isinstance(text_slots, list):
        errors.append("constraints.text_slots must be a list")
        text_slots = []
    reference_keys = {
        (str(ref.get("name", "")), int(ref.get("occurrence", 1)))
        for ref in shape_refs
        if isinstance(ref, dict)
    }
    for slot in text_slots:
        key = (str(slot.get("shape", "")), int(slot.get("occurrence", 1)))
        if key not in reference_keys:
            errors.append(f"text slot {key} is not included in shape_refs")
        minimum = int(slot.get("min_chars", -1))
        maximum = int(slot.get("max_chars", -1))
        if minimum < 0 or maximum < minimum:
            errors.append(f"text slot {key} has invalid min/max capacity")
        if "line_count" in slot and int(slot["line_count"]) < 1:
            errors.append(f"text slot {key} has invalid line_count")

    family_summary: dict[str, Any] | None = None
    if args.family:
        family_path = resolve_repo_path(args.family, repo)
        family = read_json(family_path)
        family_id = str(family.get("family_id", ""))
        compatible = [str(value) for value in module.get("compatible_master_families", [])]
        if family_id not in compatible:
            errors.append(f"module is not registered as compatible with {family_id}")
        expected_size = family.get("slide_size_emu", {})
        actual_size = {"width": int(deck.slide_width), "height": int(deck.slide_height)}
        if actual_size != expected_size:
            errors.append(
                f"module source slide size {actual_size} differs from family {expected_size}"
            )
        safe_top = int(family.get("shell_contract", {}).get("content_top_emu", -1))
        edge_tolerance_emu = 2
        if actual_bounds and (
            actual_bounds["top"] < safe_top
            or actual_bounds["left"] < -edge_tolerance_emu
            or actual_bounds["right"] > actual_size["width"] + edge_tolerance_emu
            or actual_bounds["bottom"] > actual_size["height"] + edge_tolerance_emu
        ):
            errors.append("module bounds fall outside the family content safe region")
        family_summary = {
            "family_id": family_id,
            "contract": relative_or_absolute(family_path, repo),
            "content_top_emu": safe_top,
        }

    result = {
        "ok": not errors,
        "module": relative_or_absolute(module_path, repo),
        "module_id": module.get("module_id"),
        "source": relative_or_absolute(source_deck, repo),
        "source_slide": source_slide_number,
        "source_master_sha256": actual_source_master_sha or None,
        "shape_count": len(records),
        "bounds_emu": actual_bounds,
        "geometry_sha256": actual_geometry_sha,
        "family": family_summary,
        "errors": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(2)


def cmd_render(args: argparse.Namespace) -> None:
    repo = find_repo_root(Path(args.repo_root) if args.repo_root else None)
    input_path = resolve_repo_path(args.input, repo)
    output_dir = resolve_repo_path(args.output_dir, repo)
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ppt-render-") as temporary:
        temporary_dir = Path(temporary)
        with powerpoint_app() as app:
            presentation = app.Presentations.Open(str(input_path), True, False, False)
            try:
                expected_count = int(presentation.Slides.Count)
                presentation.Export(str(temporary_dir), "PNG", args.width, args.height)
            finally:
                presentation.Close()
        rendered = [
            path for path in temporary_dir.iterdir() if path.suffix.lower() == ".png"
        ]
        def render_order(path: Path) -> tuple[int, int | str]:
            match = re.search(r"(\d+)$", path.stem)
            return (0, int(match.group(1))) if match else (1, path.name)

        rendered.sort(key=render_order)
        if len(rendered) != expected_count:
            fail(f"PowerPoint rendered {len(rendered)} PNG files; expected {expected_count}")
        expected_names = set()
        for number, source in enumerate(rendered, 1):
            destination = output_dir / f"slide-{number:03d}.png"
            shutil.copy2(source, destination)
            expected_names.add(destination.name)
        for stale in output_dir.glob("slide-*.png"):
            if stale.name not in expected_names and re.fullmatch(
                r"slide-\d{3}\.png", stale.name
            ):
                stale.unlink()
    print(output_dir)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Copy, compose, replace, inspect, verify, and render template-derived PPTX files."
    )
    parser.add_argument("--repo-root", help="Repository root; normally auto-detected")
    subparsers = parser.add_subparsers(dest="command", required=True)

    index_parser = subparsers.add_parser("index", help="Index the template library")
    index_parser.add_argument("--write", action="store_true", help="Write 模板/索引.json")
    index_parser.set_defaults(func=cmd_index)

    inspect_parser = subparsers.add_parser("inspect", help="List slides and exact shape names")
    inspect_parser.add_argument("--input", required=True)
    inspect_parser.add_argument("--output")
    inspect_parser.set_defaults(func=cmd_inspect)

    extract_parser = subparsers.add_parser("extract", help="Copy source slides into one-slide PPTX files")
    extract_parser.add_argument("--source", required=True)
    extract_parser.add_argument("--slides", type=int, nargs="+", required=True)
    extract_parser.add_argument("--category", choices=CATEGORIES)
    extract_parser.add_argument("--output-dir")
    extract_parser.add_argument("--names", nargs="+")
    extract_parser.add_argument("--overwrite", action="store_true")
    extract_parser.set_defaults(func=cmd_extract)

    learn_parser = subparsers.add_parser(
        "learn", help="Extract every source slide as a temporary curation candidate"
    )
    learn_parser.add_argument("--source-dir", default="待提取")
    learn_parser.add_argument("--overwrite", action="store_true")
    learn_parser.add_argument("--resume", action="store_true")
    learn_parser.set_defaults(func=cmd_learn)

    archive_parser = subparsers.add_parser(
        "archive-learned",
        help="Move fully learned source decks from 待提取 to 已提取",
    )
    archive_parser.set_defaults(func=cmd_archive_learned)

    compose_parser = subparsers.add_parser("compose", help="Compose an output deck entirely from template slides")
    compose_parser.add_argument("--plan", required=True)
    compose_parser.add_argument("--output", required=True)
    compose_parser.add_argument("--overwrite", action="store_true")
    compose_parser.set_defaults(func=cmd_compose)

    family_compose_parser = subparsers.add_parser(
        "compose-family",
        help="Compose a deck by duplicating role/module source pages inside one master family",
    )
    family_compose_parser.add_argument("--plan", required=True)
    family_compose_parser.add_argument("--output", required=True)
    family_compose_parser.add_argument("--overwrite", action="store_true")
    family_compose_parser.set_defaults(func=cmd_compose_family)

    text_parser = subparsers.add_parser("replace-text", help="Replace text without changing slide geometry")
    text_parser.add_argument("--input", required=True)
    text_parser.add_argument("--map", required=True)
    text_parser.add_argument("--output")
    text_parser.set_defaults(func=cmd_replace_text)

    image_parser = subparsers.add_parser("replace-image", help="Replace picture media without changing its frame")
    image_parser.add_argument("--input", required=True)
    image_parser.add_argument("--map", required=True)
    image_parser.add_argument("--output")
    image_parser.set_defaults(func=cmd_replace_image)

    verify_parser = subparsers.add_parser("verify", help="Compare output structure with its source templates")
    verify_parser.add_argument("--input", required=True)
    verify_parser.add_argument("--provenance")
    verify_parser.set_defaults(func=cmd_verify)

    layered_verify_parser = subparsers.add_parser(
        "verify-layered",
        help="Compare a master-family deck with its layered provenance sources",
    )
    layered_verify_parser.add_argument("--input", required=True)
    layered_verify_parser.add_argument("--provenance")
    layered_verify_parser.set_defaults(func=cmd_verify_layered)

    master_parser = subparsers.add_parser(
        "verify-master",
        help="Require one actual PowerPoint master and optionally enforce a master-family contract",
    )
    master_parser.add_argument("--input", required=True)
    master_parser.add_argument("--family")
    master_parser.set_defaults(func=cmd_verify_master)

    module_parser = subparsers.add_parser(
        "verify-module",
        help="Validate a below-title content module against its source and master family",
    )
    module_parser.add_argument("--module", required=True)
    module_parser.add_argument("--family")
    module_parser.set_defaults(func=cmd_verify_module)

    render_parser = subparsers.add_parser("render", help="Render every slide to PNG through PowerPoint")
    render_parser.add_argument("--input", required=True)
    render_parser.add_argument("--output-dir", required=True)
    render_parser.add_argument("--width", type=int, default=1600)
    render_parser.add_argument("--height", type=int, default=900)
    render_parser.set_defaults(func=cmd_render)

    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
